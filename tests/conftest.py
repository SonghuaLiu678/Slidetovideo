from pathlib import Path

import pymupdf
import pytest
from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.util import Inches

from slides2video.media import run_ffmpeg


@pytest.fixture()
def sample_deck(tmp_path: Path) -> Path:
    """Three slides: bullets, a table, and a big picture."""
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(10), Inches(7.5)

    s1 = prs.slides.add_slide(prs.slide_layouts[1])
    s1.shapes.title.text = "How it works"
    tf = s1.placeholders[1].text_frame
    tf.text = "First we do the first thing."
    for line in ("Then the second thing, worth 30 percent.", "Finally the third thing."):
        tf.add_paragraph().text = line

    s2 = prs.slides.add_slide(prs.slide_layouts[5])
    s2.shapes.title.text = "Earnings"
    table = s2.shapes.add_table(4, 2, Inches(1), Inches(2), Inches(8), Inches(3)).table
    table.cell(0, 0).text, table.cell(0, 1).text = "Case", "Tokens"
    for r, (a, b) in enumerate([("A", "110"), ("B", "30"), ("C", "10")], start=1):
        table.cell(r, 0).text, table.cell(r, 1).text = a, b

    s3 = prs.slides.add_slide(prs.slide_layouts[5])
    s3.shapes.title.text = "The screen"
    img = tmp_path / "shot.png"
    im = Image.new("RGB", (900, 600), "white")
    ImageDraw.Draw(im).rectangle((20, 20, 880, 120), fill="green")
    im.save(img)
    s3.shapes.add_picture(str(img), Inches(1), Inches(1.5), width=Inches(8))

    path = tmp_path / "sample.pptx"
    prs.save(str(path))
    return path


def make_tone(path: Path, seconds: float) -> None:
    run_ffmpeg(["-f", "lavfi", "-i", f"sine=frequency=440:duration={seconds}",
                "-c:a", "libmp3lame", "-b:a", "48k", str(path)])


def make_frames(out_dir: Path, n: int, size=(640, 480)) -> None:
    out_dir.mkdir(parents=True, exist_ok=True)
    for i in range(1, n + 1):
        shade = 30 + (i * 40) % 200
        Image.new("RGB", size, (shade, 255 - shade, 128)).save(out_dir / f"{i:03d}.png")


def make_pdf(path: Path, pages: int) -> None:
    doc = pymupdf.open()
    for i in range(pages):
        page = doc.new_page(width=720, height=540)
        page.insert_text((72, 100), f"page {i + 1}", fontsize=32)
    doc.save(str(path))
