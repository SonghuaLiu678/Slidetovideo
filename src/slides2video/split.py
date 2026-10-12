"""Split a deck into reveal steps.

A narrated video usually shows a slide's text piece by piece while the voice
reads it. PowerPoint does that with click animations, which a PDF export
flattens. This module does it on the slide level instead: one new slide per
step, with the not-yet-revealed text made fully transparent. Layout never
shifts, and a plain PDF export of the result has exactly one page per step.

Rules (deliberately simple, edit steps.json or the deck if you disagree):
  * the text box with the most non-empty paragraphs (at least two) is revealed
    one paragraph at a time;
  * otherwise the first table is revealed one row at a time (header stays);
  * everything else is visible from the first step;
  * slides listed in ``keep`` are not split.
"""
from __future__ import annotations

from copy import deepcopy
from pathlib import Path
from typing import Iterable, List, Optional, Sequence, Set, Tuple

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.oxml.ns import qn
from lxml import etree

from .steps import Project, Step

A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
_FILLS = ("a:noFill", "a:solidFill", "a:gradFill", "a:blipFill", "a:pattFill", "a:grpFill")
_BULLETS = ("a:buNone", "a:buAutoNum", "a:buChar", "a:buBlip")
_AFTER_BULLET = ("a:tabLst", "a:defRPr", "a:extLst")
SCREEN_MIN_AREA = 0.15  # a picture covering this share of the slide counts as a "screen"


# ---------------------------------------------------------------- text access

def _para_text(p_el) -> str:
    """Text of one a:p. A line break (a:br, typed as Shift+Enter) becomes a
    space; dropping it glued words together ("HIGHReported score")."""
    parts = []
    for node in p_el.iter():
        if node.tag == qn("a:t"):
            parts.append(node.text or "")
        elif node.tag == qn("a:br"):
            parts.append(" ")
    return " ".join("".join(parts).split())


def _cell_text(tc_el) -> str:
    return " ".join(t for t in (_para_text(p) for p in tc_el.iter(qn("a:p"))) if t)


def _shape_type(shp):
    try:
        return shp.shape_type
    except NotImplementedError:  # some placeholder kinds have no type
        return None


def _walk(shapes) -> Iterable:
    for shp in shapes:
        yield shp
        if _shape_type(shp) == MSO_SHAPE_TYPE.GROUP:
            yield from _walk(shp.shapes)


def _title_shape(slide):
    try:
        return slide.shapes.title
    except Exception:  # pragma: no cover - layouts without a title
        return None


def slide_title(slide) -> str:
    t = _title_shape(slide)
    if t is not None and t.has_text_frame and t.text_frame.text.strip():
        return " ".join(t.text_frame.text.split())
    for shp in _walk(slide.shapes):
        if shp.has_text_frame and shp.text_frame.text.strip():
            return " ".join(shp.text_frame.text.split())
    return ""


def has_screen(slide, prs) -> bool:
    area = prs.slide_width * prs.slide_height
    for shp in _walk(slide.shapes):
        if _shape_type(shp) == MSO_SHAPE_TYPE.PICTURE and shp.width and shp.height:
            if (shp.width * shp.height) / area >= SCREEN_MIN_AREA:
                return True
    return False


def _paragraph_elements(shape) -> list:
    return list(shape.text_frame._txBody.findall(qn("a:p")))


def _table_rows(shape) -> list:
    return list(shape._element.iter(qn("a:tr")))


def _pick_reveal_target(slide) -> Tuple[str, Optional[object], List[int]]:
    """Returns (kind, shape, unit_indices). kind is 'text', 'table' or 'none'."""
    title = _title_shape(slide)
    best, best_units = None, []
    for shp in _walk(slide.shapes):
        if not shp.has_text_frame or (title is not None and shp.shape_id == title.shape_id):
            continue
        units = [i for i, p in enumerate(_paragraph_elements(shp)) if _para_text(p)]
        if len(units) >= 2 and len(units) > len(best_units):
            best, best_units = shp, units
    if best is not None:
        return "text", best, best_units
    for shp in _walk(slide.shapes):
        if getattr(shp, "has_table", False) and shp.has_table:
            rows = _table_rows(shp)
            if len(rows) >= 3:
                return "table", shp, list(range(1, len(rows)))
    return "none", None, []


def _visible_text(slide, kind, target, units, upto: int) -> List[str]:
    """All visible paragraphs on the slide when ``upto`` units are shown."""
    out: List[str] = []
    shown_limit = units[upto - 1] if units and upto else -1
    for shp in _walk(slide.shapes):
        if shp.has_text_frame:
            for i, p in enumerate(_paragraph_elements(shp)):
                if kind == "text" and target is not None and shp.shape_id == target.shape_id and i > shown_limit:
                    continue
                txt = _para_text(p)
                if txt:
                    out.append(txt)
        elif getattr(shp, "has_table", False) and shp.has_table:
            for r, tr in enumerate(_table_rows(shp)):
                if kind == "table" and target is not None and shp.shape_id == target.shape_id and r > upto:
                    continue
                cells = [_cell_text(tc) for tc in tr.findall(qn("a:tc"))]
                line = " | ".join(c for c in cells if c)
                if line:
                    out.append(line)
    return out


def _unit_new_text(slide, kind, target, units, j: int) -> List[str]:
    prev = _visible_text(slide, kind, target, units, j - 1) if j > 1 else []
    cur = _visible_text(slide, kind, target, units, j)
    remaining = list(prev)
    new = []
    for line in cur:
        if line in remaining:
            remaining.remove(line)
        else:
            new.append(line)
    return new


# ------------------------------------------------------------ hiding content

def _set_transparent(rpr) -> None:
    for tag in _FILLS:
        for el in rpr.findall(qn(tag)):
            rpr.remove(el)
    fill = etree.SubElement(rpr, qn("a:solidFill"))
    rpr.remove(fill)
    clr = etree.SubElement(fill, qn("a:srgbClr"))
    clr.set("val", "FFFFFF")
    etree.SubElement(clr, qn("a:alpha")).set("val", "0")
    ln = rpr.find(qn("a:ln"))
    rpr.insert(list(rpr).index(ln) + 1 if ln is not None else 0, fill)


def _hide_paragraph(p_el) -> None:
    for run in list(p_el.findall(qn("a:r"))) + list(p_el.findall(qn("a:fld"))):
        rpr = run.find(qn("a:rPr"))
        if rpr is None:
            rpr = etree.Element(qn("a:rPr"))
            run.insert(0, rpr)
        _set_transparent(rpr)
    ppr = p_el.find(qn("a:pPr"))
    if ppr is None:
        ppr = etree.Element(qn("a:pPr"))
        p_el.insert(0, ppr)
    for tag in _BULLETS:
        for el in ppr.findall(qn(tag)):
            ppr.remove(el)
    bu = etree.Element(qn("a:buNone"))
    for i, child in enumerate(list(ppr)):
        if child.tag in [qn(t) for t in _AFTER_BULLET]:
            ppr.insert(i, bu)
            break
    else:
        ppr.append(bu)


def _hide_target(clone, kind, target_index_path, units, upto) -> None:
    """Hide everything after the first ``upto`` units on the cloned slide."""
    shapes = list(_walk(clone.shapes))
    target = shapes[target_index_path] if target_index_path is not None else None
    if target is None:
        return
    if kind == "text":
        limit = units[upto - 1]
        for i, p in enumerate(_paragraph_elements(target)):
            if i > limit:
                _hide_paragraph(p)
    elif kind == "table":
        for r, tr in enumerate(_table_rows(target)):
            if r > upto:
                for p in tr.iter(qn("a:p")):
                    _hide_paragraph(p)


# ------------------------------------------------------------------- cloning

def _clone_slide(prs, src):
    new = prs.slides.add_slide(src.slide_layout)
    for shp in list(new.shapes):
        shp._element.getparent().remove(shp._element)
    rid_map = {}
    for rel in src.part.rels.values():
        if rel.reltype.endswith("/slideLayout") or rel.reltype.endswith("/notesSlide"):
            continue
        if rel.is_external:
            rid_map[rel.rId] = new.part.relate_to(rel.target_ref, rel.reltype, is_external=True)
        else:
            rid_map[rel.rId] = new.part.relate_to(rel.target_part, rel.reltype)
    tree = new.shapes._spTree
    for el in src.shapes._spTree.iterchildren():
        if el.tag in (qn("p:nvGrpSpPr"), qn("p:grpSpPr")):
            continue
        tree.append(deepcopy(el))
    bg = src._element.cSld.find(qn("p:bg"))
    if bg is not None:
        new._element.cSld.insert(0, deepcopy(bg))
    for el in new._element.iter():
        for attr in (qn("r:embed"), qn("r:link"), qn("r:id"), qn("r:pict")):
            v = el.get(attr)
            if v in rid_map:
                el.set(attr, rid_map[v])
    return new


# ----------------------------------------------------------------- public API

def split_deck(src: Path, dst: Path, keep: Sequence[int] = ()) -> Project:
    """Write a split copy of ``src`` to ``dst`` and return the steps skeleton."""
    prs = Presentation(str(src))
    originals = list(prs.slides)
    sld_id_lst = prs.slides._sldIdLst
    orig_ids = list(sld_id_lst)[: len(originals)]
    keep_set: Set[int] = set(keep)
    steps: List[Step] = []

    for idx, slide in enumerate(originals, start=1):
        kind, target, units = ("none", None, []) if idx in keep_set else _pick_reveal_target(slide)
        n_units = len(units) if kind != "none" else 1
        shapes = list(_walk(slide.shapes))
        t_index = shapes.index(target) if target is not None else None
        title = slide_title(slide)
        screen = has_screen(slide, prs)
        for j in range(1, n_units + 1):
            clone = _clone_slide(prs, slide)
            if kind != "none":
                _hide_target(clone, kind, t_index, units, j)
            text = _visible_text(slide, kind, target, units, j) if kind != "none" else _visible_text(slide, "none", None, [], 0)
            new = _unit_new_text(slide, kind, target, units, j) if kind != "none" else list(text)
            steps.append(Step(step=len(steps) + 1, source_slide=idx, title=title,
                              text=text, new_text=new, has_screen=screen))
    for sld in orig_ids:
        prs.part.drop_rel(sld.rId)
        sld_id_lst.remove(sld)
    Path(dst).parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(dst))
    return Project(deck=Path(dst).name, steps=steps)


def steps_from_deck(path: Path) -> Project:
    """Use a deck as it is: one step per slide. ``new_text`` is whatever is not
    already on the previous slide, so decks that were split by hand work too."""
    prs = Presentation(str(path))
    steps: List[Step] = []
    prev: List[str] = []
    for idx, slide in enumerate(prs.slides, start=1):
        text = _visible_text(slide, "none", None, [], 0)
        remaining = list(prev)
        new = []
        for line in text:
            if line in remaining:
                remaining.remove(line)
            else:
                new.append(line)
        steps.append(Step(step=idx, source_slide=idx, title=slide_title(slide), text=text,
                          new_text=new or text, has_screen=has_screen(slide, prs)))
        prev = text
    return Project(deck=Path(path).name, steps=steps)
